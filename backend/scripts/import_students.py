#!/usr/bin/env python3
"""
Bulk Student Registration Importer for Qiskit Fall Fest 2026.

This script imports student records from an Excel spreadsheet (.xlsx) into
the existing Qiskit Fall Fest 2026 registration system via the backend bulk-import API.

Usage:
    python scripts/import_students.py --dry-run
    python scripts/import_students.py --file "path/to/file.xlsx" --dry-run
    python scripts/import_students.py
    python scripts/import_students.py --file "path/to/file.xlsx"
"""

import os
import sys
import csv
import json
import argparse
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

try:
    import openpyxl
except ImportError:
    print("[ERROR] 'openpyxl' is required to read Excel files. Please install it via: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

try:
    import requests
except ImportError:
    print("[ERROR] 'requests' is required for API communication. Please install it via: pip install requests", file=sys.stderr)
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_EXCEL_FILE = SCRIPT_DIR / "Untitled spreadsheet - Copy.xlsx"
DEFAULT_OUTPUT_CSV = SCRIPT_DIR / "import-results.csv"

DEFAULT_API_URL = os.environ.get("API_URL", "http://localhost:5000")
ORGANIZER_EMAIL = os.environ.get("ORGANIZER_EMAIL", "admin@qiskitfallfest.com")
ORGANIZER_PASSWORD = os.environ.get("ORGANIZER_PASSWORD", "Admin@123")


def normalize_phone(raw_value: Any) -> str:
    """Preserves phone number as string, maintaining leading zeros and handling float serialization."""
    if raw_value is None:
        return ""
    if isinstance(raw_value, float):
        if raw_value.is_integer():
            return str(int(raw_value))
        return str(raw_value)
    if isinstance(raw_value, int):
        return str(raw_value)
    
    val = str(raw_value).strip()
    if val.endswith(".0"):
        val = val[:-2]
    return val


def normalize_email(raw_value: Any) -> str:
    """Trims whitespace and lowercases the email address."""
    if raw_value is None:
        return ""
    return str(raw_value).strip().lower()


def parse_excel_file(file_path: Path) -> List[Dict[str, Any]]:
    """
    Reads the Excel file and extracts student records.
    Only maps:
      1. Name -> name
      2. Name of your College or University -> college
      3. Contact Number (WhatsApp) -> phone
      4. Email -> email
    """
    if not file_path.is_file():
        print(f"[ERROR] Excel file not found: {file_path}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active

    header_row = [cell.value for cell in sheet[1]]
    header_map = {}
    for idx, cell_val in enumerate(header_row):
        if cell_val is None:
            continue
        cleaned = str(cell_val).strip()
        header_map[cleaned] = idx

    # Validate required headers
    name_col = None
    college_col = None
    phone_col = None
    email_col = None

    for col_name, idx in header_map.items():
        lower = col_name.lower()
        if lower == "name":
            name_col = idx
        elif "college" in lower or "university" in lower:
            college_col = idx
        elif "contact" in lower or "whatsapp" in lower or "phone" in lower or "mobile" in lower:
            phone_col = idx
        elif "email" in lower:
            email_col = idx

    if name_col is None or college_col is None or phone_col is None or email_col is None:
        print(f"[ERROR] Could not locate all required columns in Excel sheet. Found headers: {list(header_map.keys())}", file=sys.stderr)
        sys.exit(1)

    records = []
    # Data starts at row 2
    for r in range(2, sheet.max_row + 1):
        row_vals = [sheet.cell(row=r, column=c + 1).value for c in range(sheet.max_column)]
        
        # Check if row is completely empty
        if all(v is None or str(v).strip() == "" for v in row_vals):
            continue

        raw_name = row_vals[name_col] if name_col < len(row_vals) else ""
        raw_college = row_vals[college_col] if college_col < len(row_vals) else ""
        raw_phone = row_vals[phone_col] if phone_col < len(row_vals) else ""
        raw_email = row_vals[email_col] if email_col < len(row_vals) else ""

        name = str(raw_name).strip() if raw_name is not None else ""
        college = str(raw_college).strip() if raw_college is not None else ""
        phone = normalize_phone(raw_phone)
        email = normalize_email(raw_email)

        # Skip empty rows where neither name nor email is provided
        if not name and not email:
            continue

        records.append({
            "row": r,
            "name": name,
            "college": college,
            "phone": phone,
            "email": email,
        })

    return records


def get_organizer_token(api_url: str) -> str:
    """Authenticates with the backend as Organizer to obtain a JWT token."""
    login_url = f"{api_url.rstrip('/')}/api/v1/organizers/login"
    payload = {
        "email": ORGANIZER_EMAIL,
        "password": ORGANIZER_PASSWORD,
    }

    try:
        resp = requests.post(login_url, json=payload, timeout=10)
        if resp.status_code != 200:
            print(f"[ERROR] Organizer authentication failed (HTTP {resp.status_code}): {resp.text}", file=sys.stderr)
            sys.exit(1)
        data = resp.json()
        token = data.get("data", {}).get("token") or data.get("token")
        if not token:
            print(f"[ERROR] No token found in authentication response: {data}", file=sys.stderr)
            sys.exit(1)
        return token
    except Exception as e:
        print(f"[ERROR] Unable to connect to backend server at {api_url}: {e}", file=sys.stderr)
        print("Please ensure the backend server is running.", file=sys.stderr)
        sys.exit(1)


def send_batch(api_url: str, token: str, students: List[Dict[str, Any]], dry_run: bool) -> Dict[str, Any]:
    """Sends a batch of students to the bulk-import endpoint."""
    endpoint = f"{api_url.rstrip('/')}/api/v1/organizer/registrations/bulk-import"
    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
        "X-Event-Profile": "pre-qiskit",
    }
    payload = {
        "dryRun": dry_run,
        "students": students,
    }

    resp = requests.post(endpoint, json=payload, headers=headers, timeout=60)
    if resp.status_code != 200:
        raise RuntimeError(f"Bulk import request failed (HTTP {resp.status_code}): {resp.text}")
    
    return resp.json()


def write_results_csv(output_file: Path, results: List[Dict[str, Any]]) -> None:
    """Writes results to CSV formatted with specified columns."""
    fieldnames = [
        "name",
        "college",
        "phone",
        "email",
        "status",
        "registration_id",
        "email_status",
        "error",
    ]

    with open(output_file, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in results:
            writer.writerow({
                "name": r.get("name", ""),
                "college": r.get("college", ""),
                "phone": r.get("phone", ""),
                "email": r.get("email", ""),
                "status": r.get("status", ""),
                "registration_id": r.get("registrationId") or r.get("registration_id") or "",
                "email_status": r.get("emailStatus") or r.get("email_status") or "",
                "error": r.get("error") or "",
            })


def main():
    parser = argparse.ArgumentParser(description="Bulk student registration importer from Excel.")
    parser.add_argument(
        "--file",
        "-f",
        type=str,
        default=None,
        help="Path to Excel (.xlsx) file. Defaults to 'Untitled spreadsheet - Copy.xlsx' in script directory.",
    )
    parser.add_argument(
        "--dry-run",
        "-d",
        action="store_true",
        help="Simulate the import: validates records and checks existing registrations without modifying database or sending emails.",
    )
    parser.add_argument(
        "--batch-size",
        "-b",
        type=int,
        default=25,
        help="Number of student records to process per request batch. Default is 25.",
    )
    parser.add_argument(
        "--output",
        "-o",
        type=str,
        default=None,
        help="Path to save import results CSV. Defaults to 'import-results.csv' in script directory.",
    )
    parser.add_argument(
        "--api-url",
        type=str,
        default=DEFAULT_API_URL,
        help=f"Backend API URL. Defaults to {DEFAULT_API_URL}.",
    )

    args = parser.parse_args()

    # Resolve Excel file path safely using pathlib
    if args.file:
        excel_path = Path(args.file).resolve()
    else:
        excel_path = DEFAULT_EXCEL_FILE

    # Resolve output CSV path
    if args.output:
        output_path = Path(args.output).resolve()
    else:
        output_path = DEFAULT_OUTPUT_CSV

    is_dry_run = args.dry_run

    # Parse Excel
    records = parse_excel_file(excel_path)
    total_records = len(records)

    # Obtain token
    token = get_organizer_token(args.api_url)

    # Process in batches
    all_results = []
    total_valid = 0
    total_invalid = 0
    total_already_registered = 0
    total_newly_registered = 0
    total_email_sent = 0
    total_email_failed = 0

    batch_size = max(1, args.batch_size)
    for i in range(0, total_records, batch_size):
        batch = records[i:i + batch_size]
        student_payloads = [
            {
                "name": r["name"],
                "college": r["college"],
                "phone": r["phone"],
                "email": r["email"],
            }
            for r in batch
        ]

        try:
            res = send_batch(args.api_url, token, student_payloads, is_dry_run)
            data = res.get("data", {})
            summary = data.get("summary", {})
            results = data.get("results", [])

            total_valid += summary.get("validRecords", 0)
            total_invalid += summary.get("invalidRecords", 0)
            total_already_registered += summary.get("alreadyRegistered", 0)
            total_newly_registered += summary.get("newlyRegistered", 0)
            total_email_sent += summary.get("emailSent", 0)
            total_email_failed += summary.get("emailFailed", 0)

            all_results.extend(results)
        except Exception as e:
            print(f"[ERROR] Batch failed: {e}", file=sys.stderr)
            for r in batch:
                all_results.append({
                    "name": r["name"],
                    "college": r["college"],
                    "phone": r["phone"],
                    "email": r["email"],
                    "status": "FAILED",
                    "registrationId": None,
                    "emailStatus": "NOT_SENT",
                    "error": str(e),
                })
                total_invalid += 1

    # Print Summary Report
    mode_str = "DRY RUN" if is_dry_run else "REAL IMPORT"
    total_failed = sum(1 for result in all_results if result.get("status") == "FAILED")
    print("\n" + "=" * 40)
    print(f"QFF26 BULK IMPORT -- {mode_str}")
    print("=" * 40)
    print(f"Excel records:             {total_records}")
    print(f"New registrations:         {total_newly_registered}")
    print(f"Already registered:        {total_already_registered}")
    print(f"Failed:                    {total_failed}")
    print(f"Emails sent:               {total_email_sent}")
    print(f"Email failures:            {total_email_failed}")

    if is_dry_run:
        print(f"Would register:            {total_newly_registered}")
        print("\nNo database changes made.")
    else:
        # Write results CSV only when an import is actually executed
        write_results_csv(output_path, all_results)
        print(f"\nResults saved to: {output_path}")

    print("=" * 40 + "\n")


if __name__ == "__main__":
    main()
